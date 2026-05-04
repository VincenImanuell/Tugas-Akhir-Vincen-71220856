"""
Ekstrak semua hasil backtesting dari folder BT_RSI/
Output: summary_bt.xlsx dengan sheet per rumusan masalah
"""

import os, glob
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PAIR_META = {
    "IMPC-HEXA":  ("Industri",                   ["Pearson", "Spearman"]),
    "ADRO-ITMG":  ("Energi",                      ["Pearson", "Spearman"]),
    "SPMA-UNIC":  ("Barang Baku",                 ["Pearson", "Spearman"]),
    "UNIC-SPMA":  ("Barang Baku",                 ["Pearson", "Spearman"]),
    "ADHI-PTPP":  ("Infrastruktur",               ["Pearson", "Spearman"]),
    "SMDR-TMAS":  ("Transportasi & Logistik",     ["Pearson", "Spearman"]),
    "PNBN-PNLF":  ("Keuangan",                    ["Pearson"]),
    "POLL-URBN":  ("Properti",                    ["Pearson", "Spearman"]),
    "DMMX-MLPT":  ("Teknologi",                   ["Pearson", "Spearman"]),
    "AMRT-ADES":  ("Barang Konsumsi Primer",      ["Pearson", "Spearman"]),
    "MAPI-PANR":  ("Barang Konsumsi Non-Primer",  ["Pearson", "Spearman"]),
    "KAEF-INAF":  ("Kesehatan",                   ["Pearson", "Spearman"]),
    "BBNI-BMRI":  ("Keuangan",                    ["Spearman"]),
    "SOSS-SKRN":  ("Industri",                    ["Kointegrasi"]),
    "KOPI-DWGL":  ("Energi",                      ["Kointegrasi"]),
    "PICO-ALMI":  ("Barang Baku",                 ["Kointegrasi"]),
    "JAST-OASA":  ("Infrastruktur",               ["Kointegrasi"]),
    "AKSI-HELI":  ("Transportasi & Logistik",     ["Kointegrasi"]),
    "JMAS-BBRI":  ("Keuangan",                    ["Kointegrasi"]),
    "JMAS-BBNI":  ("Keuangan",                    ["Kointegrasi"]),
    "NZIA-PAMG":  ("Properti",                    ["Kointegrasi"]),
    "ATIC-DIVA":  ("Teknologi",                   ["Kointegrasi"]),
    "SKBM-MGRO":  ("Barang Konsumsi Primer",      ["Kointegrasi"]),
    "GDYR-GLOB":  ("Barang Konsumsi Non-Primer",  ["Kointegrasi"]),
    "MIKA-PRIM":  ("Kesehatan",                   ["Kointegrasi"]),
}

PEARSON_PAIRS     = ["ADHI-PTPP","ADRO-ITMG","AMRT-ADES","DMMX-MLPT","IMPC-HEXA",
                     "KAEF-INAF","MAPI-PANR","PNBN-PNLF","POLL-URBN","SMDR-TMAS","SPMA-UNIC"]
SPEARMAN_PAIRS    = ["IMPC-HEXA","ADRO-ITMG","SPMA-UNIC","ADHI-PTPP","SMDR-TMAS",
                     "BBNI-BMRI","POLL-URBN","DMMX-MLPT","AMRT-ADES","MAPI-PANR","KAEF-INAF"]
KOINTEGRASI_PAIRS = ["SOSS-SKRN","KOPI-DWGL","PICO-ALMI","JAST-OASA","AKSI-HELI",
                     "JMAS-BBRI","NZIA-PAMG","ATIC-DIVA","SKBM-MGRO","GDYR-GLOB","MIKA-PRIM"]

FOLDER = "BT_RSI"


def pct(s):
    if pd.isna(s): return np.nan
    s = str(s).strip().replace("%", "")
    try: return float(s)
    except: return np.nan


def read_summary(path):
    try:
        df = pd.read_excel(path, sheet_name="Comparison Summary", header=None)
    except Exception as e:
        print(f"  ERROR membaca {path}: {e}")
        return None

    data = {}
    for _, row in df.iterrows():
        label = str(row.iloc[0]).strip() if pd.notna(row.iloc[0]) else ""
        val_b = row.iloc[1] if len(row) > 1 else np.nan
        val_r = row.iloc[2] if len(row) > 2 else np.nan
        if label:
            data[label] = (val_b, val_r)

    def g(key, idx=0):
        v = data.get(key, (np.nan, np.nan))
        return v[idx] if isinstance(v, tuple) else v

    stem = os.path.splitext(os.path.basename(path))[0]
    meta = PAIR_META.get(stem, ("Unknown", ["Unknown"]))
    y = str(g("Stock Y (Dependen)"))
    x = str(g("Stock X (Independen)"))

    sl = str(g("Stop Loss (best)")).replace("%","").strip()
    tp = str(g("Take Profit (best)")).replace("%","").strip()
    tltp = f"-{sl}/{tp}" if sl and tp else "-"

    return {
        "Pair":            stem,
        "Sektor":          meta[0],
        "Metode":          ", ".join(meta[1]),
        "Pair Type":       str(g("Pair Type")),
        "Y":               y,
        "X":               x,
        # Base
        "Base_Return":     pct(g("Total Return (%)", 0)),
        "Base_Sharpe":     pct(g("Sharpe Ratio", 0)),
        "Base_MDD":        pct(g("Max Drawdown (%)", 0)),
        "Base_WinRate":    pct(g("Win Rate (%)", 0)),
        "Base_Volatility": pct(g("Annual Volatility (%)", 0)),
        "Base_Trades":     g("Total Trades", 0),
        "Base_TP":         g("TP Hits", 0),
        "Base_SL":         g("SL Hits", 0),
        "Base_ZExit":      g("Z-Exits", 0),
        # RSI
        "RSI_Return":      pct(g("Total Return (%)", 1)),
        "RSI_Sharpe":      pct(g("Sharpe Ratio", 1)),
        "RSI_MDD":         pct(g("Max Drawdown (%)", 1)),
        "RSI_WinRate":     pct(g("Win Rate (%)", 1)),
        "RSI_Volatility":  pct(g("Annual Volatility (%)", 1)),
        "RSI_Trades":      g("Total Trades", 1),
        "RSI_TP":          g("TP Hits", 1),
        "RSI_SL":          g("SL Hits", 1),
        "RSI_ZExit":       g("Z-Exits", 1),
        # Benchmark
        "BH_Y":            pct(g(f"Buy & Hold Y ({y})")),
        "BH_X":            pct(g(f"Buy & Hold X ({x})")),
        "Base_vs_BH_Y":    pct(g("vs B&H Y (Base)")),
        "RSI_vs_BH_Y":     pct(g("vs B&H Y (RSI)")),
        "Base_vs_BH_X":    pct(g("vs B&H X (Base)")),   # ← BARU
        "RSI_vs_BH_X":     pct(g("vs B&H X (RSI)")),    # ← BARU
        # Parameters
        "z_entry":         g("z_entry (±)"),
        "z_exit":          g("z_exit"),
        "Z_Window":        g("Z-Score Window"),
        "SL":              str(g("Stop Loss (best)")),
        "TP":              str(g("Take Profit (best)")),
        "TL_TP":           tltp,
        "RSI_Period":      g("RSI Period (best)"),
        "RSI_Threshold":   g("RSI Threshold (best)"),
    }


def read_grid_search(path):
    stem = os.path.splitext(os.path.basename(path))[0]
    results = {}
    for sheet in ["Grid Search Train Z+SLTP", "Grid Search Train RSI"]:
        try:
            df = pd.read_excel(path, sheet_name=sheet)
            top = df.iloc[[0]].copy()
            top.insert(0, "Pair", stem)
            results[sheet] = top
        except:
            results[sheet] = pd.DataFrame()
    return results


# ── STYLES ───────────────────────────────────────────────────────────────────

def make_styles():
    return {
        "hdr":   Font(name="Arial", bold=True, color="FFFFFF", size=11),
        "hdr_f": PatternFill("solid", start_color="1F4E79"),
        "sec_f": PatternFill("solid", start_color="2E75B6"),
        "lbl_f": PatternFill("solid", start_color="D6E4F0"),
        "grn_f": PatternFill("solid", start_color="E2EFDA"),
        "red_f": PatternFill("solid", start_color="FFDAD9"),
        "ctr":   Alignment(horizontal="center"),
        "bdr":   Border(left=Side(style="thin"), right=Side(style="thin"),
                        top=Side(style="thin"), bottom=Side(style="thin")),
        "fnt":   Font(name="Arial", size=11),
    }

S = make_styles()


def write_df_to_sheet(ws, df, title=None, color_col=None, higher_is_better=True):
    start_row = 1
    if title:
        ws.append([title])
        ws[ws.max_row][0].font = Font(name="Arial", bold=True, size=13, color="FFFFFF")
        ws[ws.max_row][0].fill = S["hdr_f"]
        ws.merge_cells(f"A1:{get_column_letter(len(df.columns))}1")
        ws.append([])
        start_row = 3

    ws.append(list(df.columns))
    for cell in ws[ws.max_row]:
        cell.font = S["hdr"]
        cell.fill = S["hdr_f"]
        cell.alignment = S["ctr"]
        cell.border = S["bdr"]

    for _, row_data in df.iterrows():
        ws.append(list(row_data))
        row_idx = ws.max_row
        for cell in ws[row_idx]:
            cell.font = S["fnt"]
            cell.alignment = S["ctr"]
            cell.border = S["bdr"]
        if color_col and color_col in df.columns:
            col_idx = list(df.columns).index(color_col) + 1
            val = row_data[color_col]
            try:
                val = float(str(val).replace("%", ""))
                fill = S["grn_f"] if (val > 0) == higher_is_better else S["red_f"]
                ws.cell(row_idx, col_idx).fill = fill
            except:
                pass

    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col) + 4
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(w, 30)
    ws.freeze_panes = f"A{start_row + 1}"


def agg_metrics(df, prefix):
    return {
        "Return (%)":     df[f"{prefix}_Return"].mean(),
        "Sharpe Ratio":   df[f"{prefix}_Sharpe"].mean(),
        "MDD (%)":        df[f"{prefix}_MDD"].mean(),
        "Win Rate (%)":   df[f"{prefix}_WinRate"].mean(),
        "Volatility (%)": df[f"{prefix}_Volatility"].mean(),
        "Avg Trades":     df[f"{prefix}_Trades"].mean(),
        "Positive Pairs": (df[f"{prefix}_Return"] > 0).sum(),
        "Total Pairs":    len(df),
    }


# ── PARAMETER SHEET (Tabel 4.9 / 4.10 / 4.11) ────────────────────────────────

PARAM_COLS = ["Pasangan Saham", "Z Window", "±Z Entry", "±Z Exit",
              "TL/TP (%)", "RSI Period", "RSI Threshold"]

def fill_param_sheet(ws, title, pairs, df_all):
    ws.append([title])
    ws[1][0].font = Font(name="Arial", bold=True, size=13, color="FFFFFF")
    ws[1][0].fill = S["hdr_f"]
    ws.merge_cells(f"A1:{get_column_letter(len(PARAM_COLS))}1")
    ws.append([])

    ws.append(PARAM_COLS)
    for cell in ws[3]:
        cell.font = S["hdr"]
        cell.fill = S["hdr_f"]
        cell.alignment = S["ctr"]
        cell.border = S["bdr"]

    lookup = df_all.set_index("Pair") if not df_all.empty else pd.DataFrame()

    for r_idx, pair in enumerate(pairs, 4):
        if not lookup.empty and pair in lookup.index:
            row = lookup.loc[pair]
            vals = [pair, row["Z_Window"], row["z_entry"], row["z_exit"],
                    row["TL_TP"], row["RSI_Period"], row["RSI_Threshold"]]
        else:
            vals = [pair, "-", "-", "-", "-", "-", "-"]

        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = S["fnt"]
            cell.alignment = S["ctr"]
            cell.border = S["bdr"]
            if r_idx % 2 == 0:
                cell.fill = PatternFill("solid", start_color="DCE6F1")

    col_widths = [18, 10, 10, 10, 12, 12, 16]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ── MAIN ─────────────────────────────────────────────────────────────────────

files = sorted(glob.glob(os.path.join(FOLDER, "*.xlsx")))
print(f"Ditemukan {len(files)} file di {FOLDER}/\n")

rows, gs_base_all, gs_rsi_all = [], [], []

for f in files:
    stem = os.path.splitext(os.path.basename(f))[0]
    print(f"  Membaca: {stem}")
    r = read_summary(f)
    if r: rows.append(r)
    gs = read_grid_search(f)
    if not gs["Grid Search Train Z+SLTP"].empty:
        gs_base_all.append(gs["Grid Search Train Z+SLTP"])
    if not gs["Grid Search Train RSI"].empty:
        gs_rsi_all.append(gs["Grid Search Train RSI"])

df_all     = pd.DataFrame(rows)
df_gs_base = pd.concat(gs_base_all, ignore_index=True) if gs_base_all else pd.DataFrame()
df_gs_rsi  = pd.concat(gs_rsi_all,  ignore_index=True) if gs_rsi_all  else pd.DataFrame()
print(f"\nTotal pair berhasil dibaca: {len(df_all)}")

wb = Workbook()

# ─── Raw Data ────────────────────────────────────────────────────────────────
ws0 = wb.active
ws0.title = "Raw Data"
write_df_to_sheet(ws0, df_all, title="Semua Pair — Data Lengkap")

# ─── RQ1 Basis ───────────────────────────────────────────────────────────────
ws1 = wb.create_sheet("RQ1 - Per Metode (Base)")
METODE_LIST = ["Pearson", "Spearman", "Kointegrasi"]

rows_sum = []
for m in METODE_LIST:
    subset = df_all[df_all["Metode"].str.contains(m, na=False)]
    if len(subset) == 0: continue
    ag = agg_metrics(subset, "Base")
    rows_sum.append({"Metode": m, "Jumlah Pair": len(subset), **ag})

df_sum = pd.DataFrame(rows_sum)[["Metode","Jumlah Pair","Return (%)","Sharpe Ratio",
                                   "MDD (%)","Win Rate (%)","Volatility (%)","Avg Trades","Positive Pairs"]]
write_df_to_sheet(ws1, df_sum, title="RQ1 — Perbandingan Antar Metode (Basis)", color_col="Return (%)")

# ─── RQ1 RSI ─────────────────────────────────────────────────────────────────
ws1r = wb.create_sheet("RQ1 - Per Metode (RSI)")
rows_sumr = []
for m in METODE_LIST:
    subset = df_all[df_all["Metode"].str.contains(m, na=False)]
    if len(subset) == 0: continue
    ag = agg_metrics(subset, "RSI")
    rows_sumr.append({"Metode": m, "Jumlah Pair": len(subset), **ag})

df_sumr = pd.DataFrame(rows_sumr)[["Metode","Jumlah Pair","Return (%)","Sharpe Ratio",
                                     "MDD (%)","Win Rate (%)","Volatility (%)","Avg Trades","Positive Pairs"]]
write_df_to_sheet(ws1r, df_sumr, title="RQ1 — Perbandingan Antar Metode (RSI)", color_col="Return (%)")

# ─── RQ2 Per Sektor ──────────────────────────────────────────────────────────
ws2 = wb.create_sheet("RQ2 - Per Sektor")
SEKTOR_LIST = sorted(df_all["Sektor"].unique())

rows_rq2 = []
for s in SEKTOR_LIST:
    subset = df_all[df_all["Sektor"] == s]
    for prefix, label in [("Base", "Basis"), ("RSI", "RSI")]:
        ag = agg_metrics(subset, prefix)
        rows_rq2.append({
            "Sektor": s, "Strategi": label, "Jumlah Pair": ag["Total Pairs"],
            "Return (%)": round(ag["Return (%)"], 2),
            "Sharpe Ratio": round(ag["Sharpe Ratio"], 4),
            "MDD (%)": round(ag["MDD (%)"], 2),
            "Win Rate (%)": round(ag["Win Rate (%)"], 2),
            "Avg Trades": round(ag["Avg Trades"], 1),
            "Positive Pairs": int(ag["Positive Pairs"]),
        })

write_df_to_sheet(ws2, pd.DataFrame(rows_rq2),
                  title="RQ2 — Kinerja Per Sektor", color_col="Return (%)")

# ─── RQ3 Basis vs RSI ────────────────────────────────────────────────────────
ws3 = wb.create_sheet("RQ3 - Basis vs RSI")
df_rq3 = df_all[["Pair","Sektor","Metode",
                   "Base_Return","RSI_Return",
                   "Base_Sharpe","RSI_Sharpe",
                   "Base_MDD","RSI_MDD",
                   "Base_WinRate","RSI_WinRate",
                   "Base_Trades","RSI_Trades"]].copy()
df_rq3["Δ_Return"]  = df_rq3["RSI_Return"]  - df_rq3["Base_Return"]
df_rq3["Δ_Sharpe"]  = df_rq3["RSI_Sharpe"]  - df_rq3["Base_Sharpe"]
df_rq3["Δ_MDD"]     = df_rq3["RSI_MDD"]     - df_rq3["Base_MDD"]
df_rq3["Δ_WinRate"] = df_rq3["RSI_WinRate"] - df_rq3["Base_WinRate"]
df_rq3["Δ_Trades"]  = df_rq3["RSI_Trades"]  - df_rq3["Base_Trades"]
df_rq3 = df_rq3[["Pair","Sektor","Metode",
                   "Base_Return","RSI_Return","Δ_Return",
                   "Base_Sharpe","RSI_Sharpe","Δ_Sharpe",
                   "Base_MDD","RSI_MDD","Δ_MDD",
                   "Base_WinRate","RSI_WinRate","Δ_WinRate",
                   "Base_Trades","RSI_Trades","Δ_Trades"]]
write_df_to_sheet(ws3, df_rq3, title="RQ3 — Basis vs RSI-Filtered", color_col="Δ_Return")

ws3.append([])
ws3.append(["Ringkasan Agregat"])
ws3[ws3.max_row][0].font = Font(name="Arial", bold=True, color="FFFFFF")
ws3[ws3.max_row][0].fill = S["sec_f"]
for k, v in {
    "Avg Return Basis (%)":    df_all["Base_Return"].mean(),
    "Avg Return RSI (%)":      df_all["RSI_Return"].mean(),
    "Avg Sharpe Basis":        df_all["Base_Sharpe"].mean(),
    "Avg Sharpe RSI":          df_all["RSI_Sharpe"].mean(),
    "Avg MDD Basis (%)":       df_all["Base_MDD"].mean(),
    "Avg MDD RSI (%)":         df_all["RSI_MDD"].mean(),
    "Avg Win Rate Basis (%)":  df_all["Base_WinRate"].mean(),
    "Avg Win Rate RSI (%)":    df_all["RSI_WinRate"].mean(),
    "RSI lebih baik (Return)": (df_all["RSI_Return"] > df_all["Base_Return"]).sum(),
    "RSI lebih baik (MDD)":    (df_all["RSI_MDD"] > df_all["Base_MDD"]).sum(),
    "RSI lebih baik (Sharpe)": (df_all["RSI_Sharpe"] > df_all["Base_Sharpe"]).sum(),
}.items():
    ws3.append([k, round(v, 4) if isinstance(v, float) else v])

# ─── RQ4 Kinerja Keseluruhan ──────────────────────────────────────────────────
ws4 = wb.create_sheet("RQ4 - Kinerja Keseluruhan")
df_rq4 = df_all[["Pair","Sektor","Metode","Y","X",
                   "Base_Return","Base_Sharpe","Base_MDD","Base_WinRate","Base_Trades",
                   "RSI_Return","RSI_Sharpe","RSI_MDD","RSI_WinRate","RSI_Trades",
                   "BH_Y","BH_X",
                   "Base_vs_BH_Y","Base_vs_BH_X",   # ← BARU
                   "RSI_vs_BH_Y","RSI_vs_BH_X"]].copy()    # ← BARU
df_rq4.columns = ["Pair","Sektor","Metode","Y","X",
                   "Base_Return(%)","Base_Sharpe","Base_MDD(%)","Base_WinRate(%)","Base_Trades",
                   "RSI_Return(%)","RSI_Sharpe","RSI_MDD(%)","RSI_WinRate(%)","RSI_Trades",
                   "BH_Y(%)","BH_X(%)",
                   "Base_vs_BH_Y(%)","Base_vs_BH_X(%)",
                   "RSI_vs_BH_Y(%)","RSI_vs_BH_X(%)"]
write_df_to_sheet(ws4, df_rq4, title="RQ4 — Kinerja Keseluruhan Out-of-Sample", color_col="RSI_Return(%)")

ws4.append([])
ws4.append(["Statistik Keseluruhan"])
ws4[ws4.max_row][0].font = Font(name="Arial", bold=True, color="FFFFFF")
ws4[ws4.max_row][0].fill = S["sec_f"]
for r in [
    ["Total Pair",                  len(df_all)],
    ["Pair Basis Positif",          (df_all["Base_Return"] > 0).sum()],
    ["Pair RSI Positif",            (df_all["RSI_Return"] > 0).sum()],
    ["Avg Return Basis (%)",        round(df_all["Base_Return"].mean(), 2)],
    ["Avg Return RSI (%)",          round(df_all["RSI_Return"].mean(), 2)],
    ["Median Return Basis (%)",     round(df_all["Base_Return"].median(), 2)],
    ["Median Return RSI (%)",       round(df_all["RSI_Return"].median(), 2)],
    ["Avg Sharpe Basis",            round(df_all["Base_Sharpe"].mean(), 4)],
    ["Avg Sharpe RSI",              round(df_all["RSI_Sharpe"].mean(), 4)],
    ["Avg MDD Basis (%)",           round(df_all["Base_MDD"].mean(), 2)],
    ["Avg MDD RSI (%)",             round(df_all["RSI_MDD"].mean(), 2)],
    ["Avg Win Rate Basis (%)",      round(df_all["Base_WinRate"].mean(), 2)],
    ["Avg Win Rate RSI (%)",        round(df_all["RSI_WinRate"].mean(), 2)],
    ["Avg Trades Basis",            round(df_all["Base_Trades"].mean(), 1)],
    ["Avg Trades RSI",              round(df_all["RSI_Trades"].mean(), 1)],
    ["Best Pair (RSI Return)",      df_all.loc[df_all["RSI_Return"].idxmax(), "Pair"]],
    ["Worst Pair (RSI Return)",     df_all.loc[df_all["RSI_Return"].idxmin(), "Pair"]],
]:
    ws4.append(r)
    for cell in ws4[ws4.max_row]:
        cell.font = S["fnt"]
        cell.border = S["bdr"]

# ─── Parameter Tabel 4.9 / 4.10 / 4.11 ──────────────────────────────────────
ws_p1 = wb.create_sheet("Param - Pearson")
ws_p2 = wb.create_sheet("Param - Spearman")
ws_p3 = wb.create_sheet("Param - Kointegrasi")
fill_param_sheet(ws_p1, "Tabel 4.9 — Parameter Optimasi (Pearson)",     PEARSON_PAIRS,     df_all)
fill_param_sheet(ws_p2, "Tabel 4.10 — Parameter Optimasi (Spearman)",   SPEARMAN_PAIRS,    df_all)
fill_param_sheet(ws_p3, "Tabel 4.11 — Parameter Optimasi (Kointegrasi)",KOINTEGRASI_PAIRS, df_all)

# ─── Grid Search ─────────────────────────────────────────────────────────────
if not df_gs_base.empty:
    ws5 = wb.create_sheet("GS Fase1 - Terpilih (Train)")
    write_df_to_sheet(ws5, df_gs_base, title="Grid Search Fase 1 — Rank 1 per Pair (Training)")

if not df_gs_rsi.empty:
    ws6 = wb.create_sheet("GS Fase2 RSI - Terpilih (Train)")
    write_df_to_sheet(ws6, df_gs_rsi, title="Grid Search Fase 2 RSI — Rank 1 per Pair (Training)")

# ─── Simpan ──────────────────────────────────────────────────────────────────
out_path = "summary_bt.xlsx"
wb.save(out_path)
print(f"\nOutput tersimpan: {out_path}")
print("Sheet: Raw Data | RQ1 Basis | RQ1 RSI | RQ2 Sektor | RQ3 Basis vs RSI | RQ4 Keseluruhan | Param Pearson/Spearman/Kointegrasi | GS Fase1 | GS Fase2")